# -*- coding: utf-8 -*-
"""Parser das abas MH 6 / MH 12 da PEONA_INPUT (somente leitura)."""
from __future__ import annotations

from datetime import datetime, date
from openpyxl import load_workbook


def _iso(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return None


def _num(v):
    if v is None or v == "" or v == "-":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _header_lags(ws, row, start_col, n_lags):
    """Lê cabeçalhos 0,1,...,(n-1)+ da planilha."""
    cols = []
    for i in range(n_lags):
        v = ws.cell(row, start_col + i).value
        if v is None:
            cols.append(str(i) if i < n_lags - 1 else f"{i}+")
        else:
            cols.append(str(v))
    return cols


def _ler_bloco_triangulo(ws, start_row, end_row, comp_col, lag_start, n_lags, total_col):
    lag_names = _header_lags(ws, 5, lag_start, n_lags)
    rows = []
    for r in range(start_row, end_row + 1):
        comp = _iso(ws.cell(r, comp_col).value)
        if not comp:
            break
        lags = []
        for i in range(n_lags):
            cell = ws.cell(r, lag_start + i)
            raw = cell.value
            lags.append(
                {
                    "col": lag_names[i],
                    "valor": _num(raw),
                    "vazio": raw is None,
                }
            )
        rows.append(
            {
                "competencia": comp,
                "competencia_label": f"{comp[5:7]}/{comp[0:4]}",
                "lags": lags,
                "total": _num(ws.cell(r, total_col).value),
                "linha_planilha": r,
            }
        )
    return {"colunas": lag_names, "linhas": rows}


def _ler_proporcoes(ws, start_row, lag_start, n_lags, total_col):
    lag_names = _header_lags(ws, start_row, lag_start, n_lags)
    rows = []
    r = start_row + 1
    while True:
        comp = _iso(ws.cell(r, 2).value)
        if not comp:
            break
        props = []
        for i in range(n_lags):
            props.append(
                {
                    "col": lag_names[i],
                    "valor": _num(ws.cell(r, lag_start + i).value),
                }
            )
        rows.append(
            {
                "competencia": comp,
                "proporcoes": props,
                "total": _num(ws.cell(r, total_col).value),
                "linha_planilha": r,
            }
        )
        r += 1
    return {"colunas": lag_names, "linhas": rows}


def _ler_fatores(ws, start_row, lag_start, n_lags, total_col):
    return _ler_proporcoes(ws, start_row, lag_start, n_lags, total_col)


def _ler_peona_serie(ws, start_row, media_label):
    rows = []
    r = start_row
    while True:
        comp = _iso(ws.cell(r, 2).value)
        if not comp:
            break
        rows.append(
            {
                "competencia": comp,
                "total_avisado": _num(ws.cell(r, 3).value),
                "media_nm": _num(ws.cell(r, 4).value),  # Bk
                "fator_peona": _num(ws.cell(r, 5).value),  # Ik
                "peona": _num(ws.cell(r, 6).value),
                "variacao_r": _num(ws.cell(r, 7).value),
                "variacao_pct": _num(ws.cell(r, 8).value),
                "linha_planilha": r,
            }
        )
        r += 1
    return rows


def ler_aba_triangulo(wb, aba: str, janela: int, desp_label: str = "Despesa MH") -> dict:
    """
    Lê estrutura completa das abas MH/Odonto 6 ou 12.
    Layout confirmado:
      - Triângulo aviso: col B + lags + Total
      - Triângulo ocorrência: bloco à direita
      - Proporção média ~L45
      - Fator médio ~L74
      - PEONA ~L103
    """
    if aba not in wb.sheetnames:
        return {"erro": f"Aba {aba} não localizada", "aba": aba, "janela": janela}

    ws = wb[aba]
    if janela == 6:
        n_lags = 6
        aviso_lag0, aviso_total = 3, 9
        ocorr_comp, ocorr_lag0, ocorr_total = 11, 12, 18
        prop_lag0, prop_total = 3, 9
        formulas = {
            "celula_triangulo": f"SUMIFS(Base_Dados[{desp_label}]; atraso; competência de aviso)",
            "total_avisado": "SUM(lags da linha)",
            "proporcao_Mk": "SUM(coluna j na janela) / SUM(totais na janela)",
            "fator_Ik": "SUM(j × Mk_j)  — coluna Total do bloco Fator Médio",
            "base_Bk": f"AVERAGE(totais da janela de {janela} meses)",
            "peona": "Ik × Bk",
        }
        interpretacao = {
            "eixo": "aviso",
            "texto_j": (
                "Cada coluna indica há quantos meses o evento ocorreu "
                "em relação ao mês de aviso. Coluna 0 = ocorreu no próprio mês do aviso; "
                "coluna 1 = ocorreu 1 mês antes do aviso; coluna 5+ agrupa atrasos ≥ 5 meses."
            ),
        }
    else:
        n_lags = 12
        aviso_lag0, aviso_total = 3, 15
        ocorr_comp, ocorr_lag0, ocorr_total = 17, 18, 30
        prop_lag0, prop_total = 3, 15
        formulas = {
            "celula_triangulo": f"SUMIFS(Base_Dados[{desp_label}]; atraso; competência de aviso)",
            "total_avisado": "SUM(lags da linha)",
            "proporcao_Mk": "SUM(coluna j na janela) / SUM(totais na janela)",
            "fator_Ik": "SUM(j × Mk_j)  — coluna Total do bloco Fator Médio",
            "base_Bk": f"AVERAGE(totais da janela de {janela} meses)",
            "peona": "Ik × Bk",
        }
        interpretacao = {
            "eixo": "aviso",
            "texto_j": (
                "Cada coluna indica há quantos meses o evento ocorreu "
                "em relação ao mês de aviso. Coluna 0 = ocorreu no próprio mês do aviso; "
                "coluna 11+ agrupa atrasos ≥ 11 meses."
            ),
        }

    # Localizar fim do triângulo (primeira linha vazia após L6)
    fim_tri = 6
    while _iso(ws.cell(fim_tri, 2).value):
        fim_tri += 1
    fim_tri -= 1

    tri_aviso = _ler_bloco_triangulo(ws, 6, fim_tri, 2, aviso_lag0, n_lags, aviso_total)
    tri_ocorr = _ler_bloco_triangulo(
        ws, 6, fim_tri, ocorr_comp, ocorr_lag0, n_lags, ocorr_total
    )

    # Proporção: procurar rótulo
    prop_header = 45
    for r in range(40, 50):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and "Compet" in v:
            prop_header = r
            break
    proporcoes = _ler_proporcoes(ws, prop_header, prop_lag0, n_lags, prop_total)

    fat_header = 74
    for r in range(70, 80):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and "Compet" in v:
            fat_header = r
            break
    fatores = _ler_fatores(ws, fat_header, prop_lag0, n_lags, prop_total)

    peona_start = 104
    for r in range(100, 110):
        if _iso(ws.cell(r, 2).value):
            peona_start = r
            break
    peona_serie = _ler_peona_serie(ws, peona_start, f"Média {janela}m")

    # Run-off auxiliar (colunas à direita do triângulo de ocorrência)
    runoff_col = aviso_total + 2 + n_lags + 2  # approx
    # MH6: cols 20-23; MH12: 32-35
    if janela == 6:
        ro_comp, ro_puro, ro_av, ro_oc = 20, 21, 22, 23
    else:
        ro_comp, ro_puro, ro_av, ro_oc = 32, 33, 34, 35
    runoff = []
    for r in range(6, fim_tri + 1):
        comp = _iso(ws.cell(r, ro_comp).value) or _iso(ws.cell(r, 2).value)
        if not comp:
            continue
        runoff.append(
            {
                "competencia": comp,
                "runoff_puro": _num(ws.cell(r, ro_puro).value),
                "avisado_mes": _num(ws.cell(r, ro_av).value),
                "ocorrido_mes": _num(ws.cell(r, ro_oc).value),
                "linha_planilha": r,
            }
        )

    return {
        "aba": aba,
        "janela": janela,
        "arquivo": "PEONA_INPUT.xlsx",
        "intervalo_triangulo_aviso": f"B5:{chr(64+aviso_total)}{fim_tri}" if janela == 6 else f"B5:O{fim_tri}",
        "n_lags": n_lags,
        "colunas_lag": tri_aviso["colunas"],
        "triangulo_aviso": tri_aviso,
        "triangulo_ocorrencia": tri_ocorr,
        "proporcoes": proporcoes,
        "fatores": fatores,
        "peona_serie": peona_serie,
        "runoff_aux": runoff,
        "formulas": formulas,
        "interpretacao": interpretacao,
        "titulo_aviso": "Abertura das Despesas por Data de Aviso",
        "titulo_ocorrencia": "Abertura das Despesas por Data de Ocorrência",
    }


def ler_aba_mh(wb, aba: str, janela: int) -> dict:
    return ler_aba_triangulo(wb, aba, janela, desp_label="Despesa MH")


def ler_triangulos_mh(caminho_xlsx) -> dict:
    wb = load_workbook(caminho_xlsx, data_only=True)
    out = {}
    for aba, janela in (("MH 6", 6), ("MH 12", 12)):
        out[str(janela)] = ler_aba_triangulo(wb, aba, janela, desp_label="Despesa MH")
    return out


def ler_triangulos_od(caminho_xlsx) -> dict:
    wb = load_workbook(caminho_xlsx, data_only=True)
    out = {}
    for aba, janela in (("Odonto 6", 6), ("Odonto 12", 12)):
        out[str(janela)] = ler_aba_triangulo(wb, aba, janela, desp_label="Despesa OD")
    return out


if __name__ == "__main__":
    from pathlib import Path

    p = Path(__file__).resolve().parents[1] / "PEONA_INPUT.xlsx"
    for nome, fn in (("MH", ler_triangulos_mh), ("OD", ler_triangulos_od)):
        data = fn(p)
        for j, d in data.items():
            u = d["peona_serie"][-1] if d.get("peona_serie") else None
            print(nome, "janela", j, "linhas", len(d.get("triangulo_aviso", {}).get("linhas", [])), "peona", u["peona"] if u else None)
