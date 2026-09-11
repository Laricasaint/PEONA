"""
Sincroniza saídas do peona_ipasgo_v2.R + metadados da planilha PEONA_INPUT
e o comparativo RUMO × IPASGO a partir de Acompanhamento das Provisões.xlsx
para o painel executivo HTML. NÃO recalcula PEONA e NÃO altera o script R.
"""
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from historico_consistencia import (
    HistoricoImmutavelError,
    imprimir_relatorio,
    montar_consistencia_do_historico,
)
from parse_mh_triangulos import ler_triangulos_mh, ler_triangulos_od

PASTA_R = Path(
    r"W:\Gerencia de Planos e Convenios\Coordenação de Inteligência de Negócios e Informações"
    r"\02. Estudos Atuariais\006. Provisões Técnicas\005. Valores em R"
)
PASTA_PROJ = Path(__file__).resolve().parents[1]
ARQUIVO_INPUT = PASTA_PROJ / "PEONA_INPUT.xlsx"
ARQUIVO_ACOMP = PASTA_PROJ / "Acompanhamento das Provisões.xlsx"
PASTA_DATA = Path(__file__).resolve().parent / "data"

ARQUIVOS_R = [
    "peona_serie_pbi.csv",
    "peona_resumo_mes_pbi.csv",
    "peona_consistencia_pbi.csv",
    "peona_triangulo_pbi.csv",
    "PEONA_Resultado_R.xlsx",
]


def _norm_header(x) -> str:
    import re
    import unicodedata

    s = "" if x is None else str(x).strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9 ]", "", s)


def _parse_num_cell(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    x = str(v).strip()
    if not x or x.upper() in {"NA", "N/A", "-", "—", "NULL"}:
        return None
    x = x.replace(" ", "")
    if "," in x and "." in x:
        x = x.replace(".", "").replace(",", ".")
    elif "," in x:
        x = x.replace(",", ".")
    try:
        return float(x)
    except ValueError:
        return None


def _parse_mes_cell(v):
    from datetime import date, datetime

    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return date(v.year, v.month, 1)
    if isinstance(v, date):
        return date(v.year, v.month, 1)
    x = str(v).strip()
    if not x or x.lower().startswith("total"):
        return None
    # Excel serial
    try:
        n = float(x.replace(",", "."))
        if 20000 < n < 80000:
            base = datetime(1899, 12, 30)
            d = base + pd.Timedelta(days=n)
            return date(d.year, d.month, 1)
    except (TypeError, ValueError):
        pass
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%Y", "%d-%m-%Y"):
        try:
            d = datetime.strptime(x, fmt)
            return date(d.year, d.month, 1)
        except ValueError:
            continue
    # "01/2026" já coberto; tenta "jan/2026"
    try:
        d = pd.to_datetime(x, dayfirst=True, errors="coerce")
        if pd.notna(d):
            return date(int(d.year), int(d.month), 1)
    except Exception:
        pass
    return None


def ler_rumo_x_ipasgo() -> tuple[list[dict], dict]:
    """Lê RUMO × IPASGO direto do Excel Acompanhamento das Provisões (não é saída do R)."""
    aba = "Garantias Ipasgo x RUMO"
    alertas: list[str] = []
    serie: list[dict] = []
    meta: dict = {
        "fonte_arquivo": ARQUIVO_ACOMP.name,
        "fonte_aba": aba,
        "formula_diferenca": "PEONA_IPASGO - PEONA_RUMO",
        "formula_variacao": "(PEONA_IPASGO / PEONA_RUMO) - 1",
        "formula_variacao_acumulada": "(sum(PEONA_IPASGO) / sum(PEONA_RUMO)) - 1",
    }

    if not ARQUIVO_ACOMP.exists():
        alertas.append(f"Arquivo não encontrado: {ARQUIVO_ACOMP.name}")
        meta["alertas"] = alertas
        meta["reconciliado"] = False
        return [], meta

    wb = load_workbook(ARQUIVO_ACOMP, data_only=True, read_only=True)
    if aba not in wb.sheetnames:
        wb.close()
        alertas.append(f"Aba '{aba}' não encontrada em {ARQUIVO_ACOMP.name}")
        meta["alertas"] = alertas
        meta["reconciliado"] = False
        return [], meta

    ws = wb[aba]
    rows_raw = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    header_row = None
    col_mes = col_rumo = col_ipasgo = col_var = None
    for i, row in enumerate(rows_raw[:30]):
        vals = [_norm_header(c) for c in row]
        idx_mes = next((j for j, v in enumerate(vals) if v in {"mes", "competencia"}), None)
        idx_rumo = next(
            (j for j, v in enumerate(vals) if "peona rumo" in v or v == "rumo"), None
        )
        idx_ipasgo = next(
            (j for j, v in enumerate(vals) if "peona ipasgo" in v or v == "ipasgo"), None
        )
        idx_var = next(
            (j for j, v in enumerate(vals) if v in {"variacao", "var"}), None
        )
        if idx_mes is not None and idx_rumo is not None and idx_ipasgo is not None:
            header_row = i
            col_mes, col_rumo, col_ipasgo, col_var = idx_mes, idx_rumo, idx_ipasgo, idx_var
            break

    if header_row is None:
        alertas.append(f"Cabeçalho não encontrado na aba '{aba}'.")
        meta["alertas"] = alertas
        meta["reconciliado"] = False
        return [], meta

    total_planilha = {"rumo": None, "ipasgo": None, "var_pct": None}
    parsed: list[dict] = []

    for i in range(header_row + 1, len(rows_raw)):
        row = rows_raw[i]
        if col_mes >= len(row):
            continue
        rotulo = "" if row[col_mes] is None else str(row[col_mes]).strip()
        if not rotulo:
            continue
        if rotulo.lower().startswith("total"):
            total_planilha["rumo"] = _parse_num_cell(
                row[col_rumo] if col_rumo < len(row) else None
            )
            total_planilha["ipasgo"] = _parse_num_cell(
                row[col_ipasgo] if col_ipasgo < len(row) else None
            )
            if col_var is not None and col_var < len(row):
                total_planilha["var_pct"] = _parse_num_cell(row[col_var])
            continue

        mes = _parse_mes_cell(row[col_mes])
        if mes is None:
            continue
        peona_rumo = _parse_num_cell(row[col_rumo] if col_rumo < len(row) else None)
        peona_ipasgo = _parse_num_cell(row[col_ipasgo] if col_ipasgo < len(row) else None)
        if peona_rumo is None and peona_ipasgo is None:
            continue

        diferenca = (
            peona_ipasgo - peona_rumo
            if peona_rumo is not None and peona_ipasgo is not None
            else None
        )
        variacao_pct = (
            peona_ipasgo / peona_rumo - 1
            if peona_rumo not in (None, 0) and peona_ipasgo is not None
            else None
        )
        if variacao_pct is None:
            comentario = "Comparação não disponível para a competência"
        elif abs(variacao_pct) < 1e-6:
            comentario = "IPASGO próximo da RUMO"
        elif variacao_pct > 0:
            comentario = f"IPASGO {abs(variacao_pct) * 100:.2f}% acima da RUMO"
        else:
            comentario = f"IPASGO {abs(variacao_pct) * 100:.2f}% abaixo da RUMO"

        parsed.append(
            {
                "competencia": mes.strftime("%Y-%m-%d"),
                "competencia_label": mes.strftime("%m/%Y"),
                "peona_rumo": peona_rumo,
                "peona_ipasgo": peona_ipasgo,
                "diferenca": diferenca,
                "variacao_pct": variacao_pct,
                "comentario": comentario,
            }
        )

    parsed.sort(key=lambda x: x["competencia"])

    # Maior |diferença| e tendência MoM no comentário
    diffs = [(k, abs(x["diferenca"])) for k, x in enumerate(parsed) if x.get("diferenca") is not None]
    if diffs:
        idx_max = max(diffs, key=lambda t: t[1])[0]
        parsed[idx_max]["comentario"] += " · Maior diferença do período"
    for k in range(1, len(parsed)):
        d0, d1 = parsed[k - 1].get("diferenca"), parsed[k].get("diferenca")
        if d0 is None or d1 is None:
            continue
        ant, atu = abs(d0), abs(d1)
        if abs(atu - ant) < 1:
            continue
        if atu < ant:
            parsed[k]["comentario"] += " · Diferença reduziu frente ao mês anterior"
        else:
            parsed[k]["comentario"] += " · Diferença aumentou frente ao mês anterior"
        if d0 != 0 and d1 != 0 and (d0 > 0) != (d1 > 0):
            parsed[k]["comentario"] += (
                " · Mudança de direção: IPASGO passou de "
                + ("abaixo" if d0 < 0 else "acima")
                + " para "
                + ("abaixo" if d1 < 0 else "acima")
                + " da RUMO"
            )

    vals_r = [x["peona_rumo"] for x in parsed if x.get("peona_rumo") is not None]
    vals_i = [x["peona_ipasgo"] for x in parsed if x.get("peona_ipasgo") is not None]
    sum_r = sum(vals_r)
    sum_i = sum(vals_i)
    dif_acum = sum_i - sum_r
    var_acum = (sum_i / sum_r - 1) if sum_r else None

    if total_planilha["rumo"] is not None and abs(total_planilha["rumo"] - sum_r) > 0.02:
        alertas.append(
            f"TOTAL RUMO planilha ({total_planilha['rumo']:.2f}) ≠ soma das competências ({sum_r:.2f})."
        )
    if total_planilha["ipasgo"] is not None and abs(total_planilha["ipasgo"] - sum_i) > 0.02:
        alertas.append(
            f"TOTAL IPASGO planilha ({total_planilha['ipasgo']:.2f}) ≠ soma das competências ({sum_i:.2f})."
        )

    meta.update(
        {
            "n_meses": len(parsed),
            "total_peona_rumo": sum_r,
            "total_peona_ipasgo": sum_i,
            "diferenca_acumulada": dif_acum,
            "variacao_acumulada": var_acum,
            "total_planilha_rumo": total_planilha["rumo"],
            "total_planilha_ipasgo": total_planilha["ipasgo"],
            "variacao_planilha_acumulada": total_planilha["var_pct"],
            "reconciliado": len(alertas) == 0 and len(parsed) > 0,
            "alertas": alertas,
            "intervalo": {
                "inicio": parsed[0]["competencia"] if parsed else None,
                "fim": parsed[-1]["competencia"] if parsed else None,
            },
        }
    )

    # Cache CSV no painel (export / referência), gerado a partir do Excel
    if parsed:
        csv_path = PASTA_DATA / "peona_rumo_x_ipasgo_pbi.csv"
        meta_path = PASTA_DATA / "peona_rumo_x_ipasgo_meta.json"
        PASTA_DATA.mkdir(parents=True, exist_ok=True)
        df = pd.DataFrame(parsed)
        df.to_csv(csv_path, sep=";", decimal=",", index=False, encoding="utf-8")
        meta_path.write_text(
            json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    return parsed, meta


def ler_csv(nome: str) -> pd.DataFrame:
    return pd.read_csv(PASTA_DATA / nome, sep=";", decimal=",", encoding="utf-8")


def df_regs(df: pd.DataFrame) -> list[dict]:
    out = df.copy()
    for col in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[col]):
            out[col] = out[col].dt.strftime("%Y-%m-%d")
        else:
            out[col] = out[col].where(pd.notna(out[col]), None)
    return json.loads(out.to_json(orient="records", force_ascii=False))


def _num(v):
    if v is None or v == "" or v == "-":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _comp_iso(v):
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d"), v.strftime("%m/%Y")
    return None, None


def ler_planilha_input() -> dict:
    """Lê abas auxiliares da PEONA_INPUT (mesma fonte do R), sem recalcular."""
    if not ARQUIVO_INPUT.exists():
        return {
            "obs": [],
            "resultado_aba": [],
            "comparativo_ipasgo_rodarte": None,
            "input_resumo": {"erro": "PEONA_INPUT.xlsx não encontrado"},
            "consistencia": [],
            "consistencia_meta": {},
            "consistencia_acompanhamento": [],
        }

    wb = load_workbook(ARQUIVO_INPUT, data_only=True)

    # OBS
    obs = []
    if "OBS" in wb.sheetnames:
        for row in wb["OBS"].iter_rows(values_only=True):
            if row and row[0]:
                obs.append(str(row[0]))

    # Resultado
    resultado_aba = []
    comparativo = None
    if "Resultado" in wb.sheetnames:
        ws = wb["Resultado"]
        rows = list(ws.iter_rows(values_only=True))
        for row in rows:
            if not row:
                continue
            # competência mensal
            if hasattr(row[1], "strftime") and row[2] is not None:
                resultado_aba.append(
                    {
                        "competencia": row[1].strftime("%Y-%m-%d"),
                        "competencia_label": row[1].strftime("%b/%Y"),
                        "peona_mh": float(row[2]),
                        "peona_od": float(row[3]),
                        "peona_total": float(row[4]),
                    }
                )
            if row[1] == "IPASGO":
                comparativo = {
                    "ipasgo_mh": float(row[2]),
                    "ipasgo_od": float(row[3]),
                    "ipasgo_total": float(row[4]),
                }
            if row[1] == "RODARTE" and comparativo:
                comparativo.update(
                    {
                        "rodarte_mh": float(row[2]),
                        "rodarte_od": float(row[3]),
                        "rodarte_total": float(row[4]),
                    }
                )
            if row[1] == "VAR" and comparativo:
                comparativo.update(
                    {
                        "var_mh": float(row[2]),
                        "var_od": float(row[3]),
                        "var_total": float(row[4]),
                    }
                )
            if isinstance(row[1], str) and "Teste de Consistência" in row[1]:
                if comparativo is None:
                    comparativo = {}
                comparativo["texto_consistencia_resultado"] = row[1]
            if isinstance(row[1], str) and "Variação em 12 meses" in row[1]:
                if comparativo is None:
                    comparativo = {}
                comparativo["variacao_12m"] = {
                    "mh": float(row[2]) if row[2] is not None else None,
                    "od": float(row[3]) if row[3] is not None else None,
                    "total": float(row[4]) if row[4] is not None else None,
                }

    # Resumo Input / Dados
    n_dados = 0
    if "Dados" in wb.sheetnames:
        ws = wb["Dados"]
        n_dados = max(ws.max_row - 2, 0)
    input_resumo = {
        "arquivo": str(ARQUIVO_INPUT.name),
        "caminho": str(ARQUIVO_INPUT),
        "abas": wb.sheetnames,
        "registros_dados_aprox": n_dados,
        "parametros_r": {
            "MES_CORTE_HISTORICO": "2025-10-01",
            "MES_CORTE_CONSISTENCIA": "2025-09-01",
            "N_JANELA_6": 6,
            "N_JANELA_12": 12,
            "MES_APURACAO": "automático (último mês do triângulo)",
            "ARQUIVO_XLSX": "PEONA_INPUT.xlsx",
        },
        "bases_consideradas": [
            "Aba Dados (base granular ocorrência × aviso)",
            "Aba MH 6 / Odonto 6 (triângulo histórico até out/2025)",
            "Aba Teste Consistência (PEONA_INPUT · 6m a partir de out/2025)",
            "Acompanhamento das Provisões (regime 12m até set/2025)",
            "Aba Garantias Ipasgo x RUMO (comparativo RUMO × IPASGO)",
        ],
    }

    triangulos_mh = ler_triangulos_mh(ARQUIVO_INPUT)
    triangulos_od = ler_triangulos_od(ARQUIVO_INPUT)

    return {
        "obs": obs,
        "resultado_aba": resultado_aba,
        "comparativo_ipasgo_rodarte": comparativo,
        "input_resumo": input_resumo,
        "consistencia": [],
        "consistencia_meta": {},
        "consistencia_acompanhamento": [],
        "consistencia_relatorio": {},
        "triangulos_mh": triangulos_mh,
        "triangulos_od": triangulos_od,
    }


def main() -> None:
    PASTA_DATA.mkdir(parents=True, exist_ok=True)
    alertas = []

    if PASTA_R.exists():
        for nome in ARQUIVOS_R:
            origem = PASTA_R / nome
            if origem.exists():
                shutil.copy2(origem, PASTA_DATA / nome)
            else:
                alertas.append(
                    {
                        "nivel": "erro",
                        "codigo": "ARQUIVO_AUSENTE",
                        "mensagem": f"Saída R ausente: {nome}",
                    }
                )
    else:
        alertas.append(
            {
                "nivel": "erro",
                "codigo": "PASTA_R_AUSENTE",
                "mensagem": f"Pasta de saída do R não encontrada: {PASTA_R}",
            }
        )

    serie = ler_csv("peona_serie_pbi.csv")
    resumo = ler_csv("peona_resumo_mes_pbi.csv")
    tri = ler_csv("peona_triangulo_pbi.csv")

    serie["competencia"] = pd.to_datetime(serie["competencia"]).dt.strftime("%Y-%m-%d")
    tri["comp_aviso"] = pd.to_datetime(tri["comp_aviso"]).dt.strftime("%Y-%m-%d")

    s6 = serie[serie["janela"] == 6].copy()
    consolidado_rows = []
    for comp, grp in s6.groupby("competencia"):
        mh = grp[grp["tipo"] == "MH"]
        od = grp[grp["tipo"] == "OD"]
        if mh.empty or od.empty:
            continue
        consolidado_rows.append(
            {
                "competencia": comp,
                "competencia_label": mh["competencia_label"].iloc[0],
                "peona_mh": float(mh["peona"].iloc[0]),
                "peona_od": float(od["peona"].iloc[0]),
                "peona_total": float(mh["peona"].iloc[0] + od["peona"].iloc[0]),
                "fator_mh": float(mh["fator_peona"].iloc[0]),
                "fator_od": float(od["fator_peona"].iloc[0]),
                "media_mh": float(mh["media_nm"].iloc[0]),
                "media_od": float(od["media_nm"].iloc[0]),
                "total_avisado_mh": float(mh["total_avisado"].iloc[0]),
                "total_avisado_od": float(od["total_avisado"].iloc[0]),
                "variacao_r_mh": None
                if pd.isna(mh["variacao_r"].iloc[0])
                else float(mh["variacao_r"].iloc[0]),
                "variacao_r_od": None
                if pd.isna(od["variacao_r"].iloc[0])
                else float(od["variacao_r"].iloc[0]),
                "variacao_pct_mh": None
                if pd.isna(mh["variacao_pct"].iloc[0])
                else float(mh["variacao_pct"].iloc[0]),
                "variacao_pct_od": None
                if pd.isna(od["variacao_pct"].iloc[0])
                else float(od["variacao_pct"].iloc[0]),
                "Mk_mh": {f"Mk{i}": float(mh[f"Mk{i}"].iloc[0]) for i in range(6)},
                "Mk_od": {f"Mk{i}": float(od[f"Mk{i}"].iloc[0]) for i in range(6)},
            }
        )

    consolidado = pd.DataFrame(consolidado_rows).sort_values("competencia")
    consolidado["variacao_r_total"] = consolidado["peona_total"].diff()
    consolidado["variacao_pct_total"] = consolidado["peona_total"].pct_change()

    planilha = ler_planilha_input()
    rumo_x_ipasgo, rumo_meta = ler_rumo_x_ipasgo()
    if not rumo_x_ipasgo:
        alertas.append(
            {
                "nivel": "aviso",
                "codigo": "RUMO_IPASGO_AUSENTE",
                "mensagem": (
                    "Série RUMO × IPASGO ausente. Verifique o arquivo "
                    f"{ARQUIVO_ACOMP.name}, aba Garantias Ipasgo x RUMO."
                ),
            }
        )
    for a in rumo_meta.get("alertas") or []:
        alertas.append(
            {
                "nivel": "aviso",
                "codigo": "RUMO_IPASGO_VALIDACAO",
                "mensagem": str(a),
            }
        )

    competencias = sorted(consolidado["competencia"].tolist())
    peona_ref = competencias[-1] if competencias else None

    # Teste de Consistência — histórico imutável (somente acréscimo)
    try:
        consist_pack = montar_consistencia_do_historico(peona_ref)
    except HistoricoImmutavelError as e:
        raise SystemExit(
            f"ERRO: atualização do painel interrompida — histórico imutável violado.\n{e}"
        ) from e

    consist_recs = consist_pack.get("registros") or []
    consist_meta = consist_pack.get("meta") or {}
    consist_acomp = consist_pack.get("acompanhamento_total") or []
    consist_rel = consist_pack.get("relatorio") or {}

    if not consist_recs:
        alertas.append(
            {
                "nivel": "erro",
                "codigo": "CONSISTENCIA_HISTORICO",
                "mensagem": (
                    "Histórico de Consistência vazio. Execute "
                    "`python historico_consistencia.py --seed`."
                ),
            }
        )
    if not ARQUIVO_ACOMP.exists():
        alertas.append(
            {
                "nivel": "aviso",
                "codigo": "ACOMPANHAMENTO_AUSENTE",
                "mensagem": f"Arquivo não encontrado: {ARQUIVO_ACOMP.name}",
            }
        )

    meta = {
        "gerado_em": datetime.now().isoformat(timespec="seconds"),
        "fonte_r": str(PASTA_R),
        "arquivo_script": "peona_ipasgo_v2.R",
        "arquivo_input": ARQUIVO_INPUT.name,
        "mes_apuracao": str(resumo["mes_apuracao"].iloc[0]) if len(resumo) else None,
        "competencia_mais_recente": peona_ref,
        "competencias_disponiveis": competencias,
        "n_competencias_6m": len(competencias),
        "janela_metodologica_principal": 6,
        "titulo_painel": "Painel Executivo de Provisões Técnicas — PEONA",
        "alertas_sync": alertas,
        "obs": planilha["obs"],
        "input_resumo": planilha["input_resumo"],
        "consistencia_relatorio": consist_rel,
    }

    payload = {
        "meta": meta,
        "resumo_mes": df_regs(resumo),
        "serie": df_regs(serie),
        "consolidado_6m": df_regs(consolidado),
        "consistencia": consist_recs,
        "consistencia_meta": consist_meta,
        "consistencia_acompanhamento": consist_acomp,
        "triangulo": df_regs(tri),
        "triangulos_mh": planilha.get("triangulos_mh") or {},
        "triangulos_od": planilha.get("triangulos_od") or {},
        "rumo_x_ipasgo": rumo_x_ipasgo,
        "rumo_x_ipasgo_meta": rumo_meta,
        "resultado_aba": planilha["resultado_aba"],
        "comparativo_ipasgo_rodarte": planilha["comparativo_ipasgo_rodarte"],
    }

    (PASTA_DATA / "peona_data.js").write_text(
        "window.PEONA_DATA = "
        + json.dumps(payload, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )
    print("OK — painel sincronizado")
    print("Competências:", len(competencias), "| Apuração:", meta["mes_apuracao"])
    print("RUMO × IPASGO:", len(rumo_x_ipasgo), "linhas")
    print("Consistência (histórico imutável):", len(consist_recs), "linhas")
    if consist_rel:
        imprimir_relatorio(consist_rel)


if __name__ == "__main__":
    main()
