# -*- coding: utf-8 -*-
"""
Teste de Consistência — leitura híbrida sem recálculo metodológico.

Regra oficial:
  · competência <= set/2025 → regime 12m (Acompanhamento das Provisões consolidado;
    detalhe MH/OD pelos valores já registrados no bloco 12m da PEONA_INPUT —
    sem usar o bloco 6m nem recalcular competências antigas).
  · competência >= out/2025 → PEONA_INPUT · bloco 6m (MH e OD).

Variação com sinal (conforme planilhas):
  V_t = PEONA_estimada / PEONA_real - 1
      = (PEONA_estimada - PEONA_real) / PEONA_real
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

CORTE_METODOLOGIA = "2025-09-01"  # até set/2025 inclusive → 12m
INICIO_6M = "2025-10-01"

# Limites ANS — operadora de grande porte
ANS_MEDIA_INF = -0.1078
ANS_MEDIA_SUP = 0.1362
ANS_MENSAL_INF = -0.2453
ANS_MENSAL_SUP = 0.3094

PASTA_PROJ = Path(__file__).resolve().parents[1]
ARQUIVO_ACOMP = PASTA_PROJ / "Acompanhamento das Provisões.xlsx"
ARQUIVO_INPUT = PASTA_PROJ / "PEONA_INPUT.xlsx"


def _num(v):
    if v is None or v == "" or v == "-":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _comp_iso(v):
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d"), v.strftime("%m/%Y")
    return None, None


def _variacao_sinal(estimada, real):
    """V = estimada/real - 1. Não converte ausência em zero."""
    if estimada is None or real is None:
        return None
    if real == 0:
        return None
    return estimada / real - 1.0


def _class_ans(valor, inf, sup):
    if valor is None:
        return {
            "codigo": "nao_avaliavel",
            "rotulo": "Não avaliável",
            "mensagem": "Valor necessário ausente na fonte oficial.",
        }
    if inf <= valor <= sup:
        return {
            "codigo": "dentro",
            "rotulo": "Dentro do parâmetro ANS",
            "mensagem": f"{valor:.2%} no intervalo [{inf:.2%}; {sup:.2%}].",
        }
    return {
        "codigo": "fora",
        "rotulo": "Fora do parâmetro ANS",
        "mensagem": f"{valor:.2%} fora de [{inf:.2%}; {sup:.2%}].",
    }


def _status_consolidado(pendente, motivo_pendente, st_mensal, st_media, comp_iso: str):
    """
    Status consolidado (grande porte) com tratamento especial de histórico:
    - se mensal é avaliável e a média não está disponível (jan/2025..set/2025),
      retorna "Avaliação parcial" (sem classificar como consistente/ não consistente).
    """
    if pendente:
        return {
            "status": "nao_executado",
            "status_rotulo": "Pendente",
            "status_mensagem": motivo_pendente
            or "Teste ainda não executado na fonte oficial.",
            "executado": False,
            "runoff_completo": False,
        }

    historico = comp_iso <= CORTE_METODOLOGIA

    # Mensal ausente => não avaliável (sem "parcial")
    if st_mensal["codigo"] == "nao_avaliavel":
        return {
            "status": "nao_avaliavel",
            "status_rotulo": "Não avaliável",
            "status_mensagem": "Falta valor necessário para aplicar os critérios ANS (mensal).",
            "executado": False,
            "runoff_completo": False,
        }

    # Média ausente no histórico => avaliação parcial
    if st_media["codigo"] in ("nao_avaliavel", "nao_disponivel"):
        if historico and st_mensal["codigo"] in ("dentro", "fora"):
            return {
                "status": "avaliacao_parcial",
                "status_rotulo": "Avaliação parcial",
                "status_mensagem": "Critério de variação média não disponível para aplicar os dois limites ANS.",
                "executado": True,
                "runoff_completo": True,
            }
        return {
            "status": "nao_avaliavel",
            "status_rotulo": "Não avaliável",
            "status_mensagem": "Falta valor necessário para aplicar os critérios ANS (média).",
            "executado": False,
            "runoff_completo": False,
        }

    if st_mensal["codigo"] == "dentro" and st_media["codigo"] == "dentro":
        return {
            "status": "consistente",
            "status_rotulo": "Consistente",
            "status_mensagem": "Os dois critérios ANS estão dentro dos limites.",
            "executado": True,
            "runoff_completo": True,
        }
    return {
        "status": "nao_consistente",
        "status_rotulo": "Não consistente",
        "status_mensagem": "Um ou mais critérios ANS estão fora dos limites.",
        "executado": True,
        "runoff_completo": True,
    }


def _meses_janela(comp_iso: str, n: int) -> list[str]:
    dt = datetime.strptime(comp_iso, "%Y-%m-%d")
    out = []
    y, m = dt.year, dt.month
    for _ in range(n):
        out.append(f"{y:04d}-{m:02d}-01")
        m -= 1
        if m == 0:
            m = 12
            y -= 1
    return list(reversed(out))


def ler_acompanhamento(caminho: Path | None = None) -> list[dict]:
    path = caminho or ARQUIVO_ACOMP
    if not path.exists():
        return []
    wb = load_workbook(path, data_only=True)
    nome = next((s for s in wb.sheetnames if "onsist" in s.lower()), None)
    if not nome:
        return []
    ws = wb[nome]
    rows = []
    for r in range(1, ws.max_row + 1):
        raw = ws.cell(r, 2).value
        if raw is None or isinstance(raw, str):
            continue
        comp, label = _comp_iso(raw)
        if not comp:
            continue
        estimada = _num(ws.cell(r, 3).value)
        real = _num(ws.cell(r, 4).value)
        var_planilha = _num(ws.cell(r, 5).value)
        var = var_planilha if var_planilha is not None else _variacao_sinal(estimada, real)
        rows.append(
            {
                "competencia": comp,
                "competencia_label": label,
                "tipo": "TOTAL",
                "peona_estimada": estimada,
                "peona_real": real,
                "variacao_mensal": var,
                "diferenca": (estimada - real)
                if estimada is not None and real is not None
                else None,
                "fonte_arquivo": path.name,
                "fonte_aba": nome,
                "fonte": f"{path.name} / {nome}",
            }
        )
    return rows


def ler_peona_input_bloco(wb, janela: str) -> list[dict]:
    nome = next((s for s in wb.sheetnames if s.startswith("Teste Consist")), None)
    if not nome:
        return []
    ws = wb[nome]
    titulo = None
    chave = "12 meses" if janela == "12m" else "6 meses"
    for r in range(1, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if isinstance(b, str) and chave in b.lower():
            titulo = r
            break
    if not titulo:
        return []

    col_hdr = titulo + 2
    out = []
    r = col_hdr + 1
    while r <= ws.max_row:
        b0 = ws.cell(r, 2).value
        if isinstance(b0, str) and "metodologia" in b0.lower():
            break
        grupos = [
            ("MH", 2, 3, 4, 5, 6, 7),
            ("OD", 9, 10, 11, 12, 13, 14),
        ]
        alguma = False
        for tipo, c_comp, c_peona, c_run, c_dif, c_dpct, c_ader in grupos:
            raw_comp = ws.cell(r, c_comp).value
            if raw_comp is None or isinstance(raw_comp, str):
                continue
            comp, label = _comp_iso(raw_comp)
            if not comp:
                continue
            estimada = _num(ws.cell(r, c_peona).value)
            real = _num(ws.cell(r, c_run).value)
            dif_plan = _num(ws.cell(r, c_dif).value)
            dpct_abs = _num(ws.cell(r, c_dpct).value)
            ader = _num(ws.cell(r, c_ader).value)
            var = _variacao_sinal(estimada, real)
            dif = (
                dif_plan
                if dif_plan is not None
                else (
                    (estimada - real)
                    if estimada is not None and real is not None
                    else None
                )
            )
            out.append(
                {
                    "competencia": comp,
                    "competencia_label": label,
                    "tipo": tipo,
                    "janela_bloco": janela,
                    "peona_estimada": estimada,
                    "peona_real": real,
                    "diferenca": dif,
                    "diferenca_pct_abs": dpct_abs,
                    "aderencia": ader,
                    "variacao_mensal": var,
                    "fonte_arquivo": ARQUIVO_INPUT.name,
                    "fonte_aba": nome,
                    "fonte": f"{ARQUIVO_INPUT.name} / {nome} · bloco {janela}",
                }
            )
            alguma = True
        if not alguma and b0 is None and ws.cell(r, 9).value is None:
            r += 1
            continue
        r += 1
    return out


def _enriquecer(serie: list[dict]) -> list[dict]:
    by_tipo: dict[str, list[dict]] = {}
    for row in serie:
        by_tipo.setdefault(row["tipo"], []).append(row)
    for tipo in by_tipo:
        by_tipo[tipo].sort(key=lambda x: x["competencia"])

    mapa = {(r["tipo"], r["competencia"], r.get("janela") or r.get("janela_bloco")): r for r in serie}
    # fallback mapa sem janela para média (mesma janela)
    out = []
    for row in serie:
        comp = row["competencia"]
        # Janela do registro: respeita bloco de origem; histórico ≤ set/2025 = 12m
        janela_bloco = row.get("janela") or row.get("janela_bloco")
        if janela_bloco in ("12m", "6m"):
            n = 12 if janela_bloco == "12m" else 6
            janela = janela_bloco
        else:
            n = 12 if comp <= CORTE_METODOLOGIA else 6
            janela = "12m" if n == 12 else "6m"
        comps_janela = _meses_janela(comp, n)
        vals = []
        incompleta = False
        for c in comps_janela:
            prev = mapa.get((row["tipo"], c, janela))
            if prev is None:
                # tenta qualquer registro do mesmo tipo/comp com mesma janela implícita
                prev = next(
                    (
                        x
                        for x in serie
                        if x["tipo"] == row["tipo"]
                        and x["competencia"] == c
                        and (x.get("janela") or x.get("janela_bloco") or ("12m" if x["competencia"] <= CORTE_METODOLOGIA else "6m"))
                        == janela
                    ),
                    None,
                )
            if prev is None or prev.get("variacao_mensal") is None:
                incompleta = True
                break
            vals.append(prev["variacao_mensal"])
        media = (sum(vals) / len(vals)) if vals and not incompleta else None

        pendente = False
        motivo = None
        real = row.get("peona_real")
        estimada = row.get("peona_estimada")
        var = row.get("variacao_mensal")
        if estimada is not None and (real is None or real == 0 or var is None):
            if real == 0 or real is None:
                pendente = True
                motivo = (
                    "Run-off puro ainda não disponível na fonte oficial "
                    "(teste não executado para a competência)."
                )
            elif var is None:
                pendente = True
                motivo = "Variação mensal não disponível na fonte oficial."

        st_m = _class_ans(None if pendente else var, ANS_MENSAL_INF, ANS_MENSAL_SUP)

        # Quando a média não está disponível, tratamos como "Não disponível" no histórico.
        if pendente:
            st_a = _class_ans(None, ANS_MEDIA_INF, ANS_MEDIA_SUP)
        else:
            if media is None:
                st_a = {
                    "codigo": "nao_disponivel",
                    "rotulo": "Não disponível",
                    "mensagem": "Variação média da janela não disponível na fonte oficial.",
                }
            else:
                st_a = _class_ans(media, ANS_MEDIA_INF, ANS_MEDIA_SUP)

        cons = _status_consolidado(pendente, motivo, st_m, st_a, comp)

        rec = {
            **row,
            "periodo_resumo": False,
            "periodo_texto": None,
            "janela": janela,
            "janela_meses": n,
            "peona": estimada,
            "runoff_puro": real,
            "diferenca_pct": var,
            "variacao_media": media,
            "media_completa": media is not None,
            "status_mensal_ans": st_m["codigo"],
            "status_mensal_rotulo": st_m["rotulo"],
            "status_mensal_mensagem": st_m["mensagem"],
            "status_media_ans": st_a["codigo"],
            "status_media_rotulo": st_a["rotulo"],
            "status_media_mensagem": st_a["mensagem"],
            "ans_mensal_inf": ANS_MENSAL_INF,
            "ans_mensal_sup": ANS_MENSAL_SUP,
            "ans_media_inf": ANS_MEDIA_INF,
            "ans_media_sup": ANS_MEDIA_SUP,
            "mudanca_metodologica": INICIO_6M,
            "pendente": pendente,
            **cons,
        }
        if pendente:
            rec["status"] = "nao_executado"
            rec["status_rotulo"] = "Pendente"
        out.append(rec)
    return out


def montar_consistencia_oficial(
    caminho_input: Path | None = None,
    caminho_acomp: Path | None = None,
) -> dict:
    path_in = caminho_input or ARQUIVO_INPUT
    path_ac = caminho_acomp or ARQUIVO_ACOMP

    acomp = ler_acompanhamento(path_ac)
    acomp_ate_set = [r for r in acomp if r["competencia"] <= CORTE_METODOLOGIA]

    input_12: list[dict] = []
    input_6: list[dict] = []
    if path_in.exists():
        wb = load_workbook(path_in, data_only=True)
        input_12 = ler_peona_input_bloco(wb, "12m")
        input_6 = ler_peona_input_bloco(wb, "6m")

    hist = []
    for r in input_12:
        if r["competencia"] <= CORTE_METODOLOGIA:
            hist.append(
                {
                    **r,
                    "regime": "historico_12m",
                    "fonte_regra": "acompanhamento_provisoes",
                    "fonte_detalhe": (
                        "Valores MH/OD do bloco 12 meses já registrados na PEONA_INPUT; "
                        "acompanhamento oficial consolidado em Acompanhamento das Provisões "
                        "(sem recálculo pela metodologia 6m)."
                    ),
                    "fonte": (
                        f"Acompanhamento das Provisões (regime 12m) · detalhe {r['tipo']} "
                        f"via {r['fonte']}"
                    ),
                }
            )

    atual = []
    for r in input_6:
        if r["competencia"] >= INICIO_6M:
            atual.append(
                {
                    **r,
                    "regime": "atual_6m",
                    "fonte_regra": "peona_input",
                    "fonte_detalhe": "PEONA_INPUT · Teste Consistência · metodologia 6 meses.",
                    "fonte": r["fonte"],
                    "janela": "6m",
                    "janela_bloco": "6m",
                }
            )

    # Pré-corte 6m: só para completar a média móvel da série 6m (não é regime oficial)
    aux_6_pre = []
    for r in input_6:
        if r["competencia"] < INICIO_6M:
            aux_6_pre.append(
                {
                    **r,
                    "regime": "aux_6m_pre_corte",
                    "fonte_regra": "peona_input",
                    "fonte_detalhe": "Bloco 6m pré-corte — apoio à média; não substitui o histórico 12m.",
                    "fonte": r["fonte"],
                    "janela": "6m",
                    "janela_bloco": "6m",
                }
            )

    # Em paralelo (a partir de out/2025): janela 12 meses — não substitui o 6m
    atual_12 = []
    for r in input_12:
        if r["competencia"] >= INICIO_6M:
            atual_12.append(
                {
                    **r,
                    "regime": "atual_12m",
                    "fonte_regra": "peona_input",
                    "fonte_detalhe": "PEONA_INPUT · Teste Consistência · metodologia 12 meses.",
                    "fonte": r["fonte"],
                    "janela": "12m",
                    "janela_bloco": "12m",
                }
            )

    registros_all = _enriquecer(hist + aux_6_pre + atual + atual_12)
    # Remove auxiliares pré-corte da série oficial exportada
    registros = [
        r for r in registros_all if r.get("regime") != "aux_6m_pre_corte"
    ]

    acomp_enriq = _enriquecer(
        [
            {
                **r,
                "regime": "historico_12m",
                "fonte_regra": "acompanhamento_provisoes",
                "fonte_detalhe": "Planilha Acompanhamento das Provisões · TESTE DE CONSISTENCIA.",
                "janela_bloco": "12m",
            }
            for r in acomp_ate_set
        ]
    )

    return {
        "registros": registros,
        "acompanhamento_total": acomp_enriq,
        "meta": {
            "corte_metodologia": CORTE_METODOLOGIA,
            "inicio_6m": INICIO_6M,
            "formula_variacao": "V_t = PEONA_estimada / PEONA_real - 1",
            "formula_variacao_latex": (
                r"V_t=\left(\frac{\operatorname{PEONA}^{\mathrm{estimada}}_t}"
                r"{\operatorname{PEONA}^{\mathrm{real}}_t}-1\right)\times 100"
            ),
            "ans_grande_porte": {
                "media": {"inf": ANS_MEDIA_INF, "sup": ANS_MEDIA_SUP},
                "mensal": {"inf": ANS_MENSAL_INF, "sup": ANS_MENSAL_SUP},
            },
            "arquivo_acompanhamento": path_ac.name if path_ac.exists() else None,
            "arquivo_input": path_in.name if path_in.exists() else None,
            "n_historico_mh_od": len(hist),
            "n_atual_mh_od": len(atual),
            "n_atual_12m_mh_od": len(atual_12),
            "n_acompanhamento": len(acomp_enriq),
            "defasagem_teste_meses": 3,
        },
    }


if __name__ == "__main__":
    p = montar_consistencia_oficial()
    print("meta", p["meta"])
    for tipo in ("MH", "OD"):
        rows = [r for r in p["registros"] if r["tipo"] == tipo]
        print(tipo, "n=", len(rows))
        for r in rows:
            if r["competencia"] in ("2025-09-01", "2025-10-01", "2026-07-01"):
                print(
                    r["competencia"],
                    r["janela"],
                    r.get("variacao_mensal"),
                    r.get("variacao_media"),
                    r["status_rotulo"],
                )
